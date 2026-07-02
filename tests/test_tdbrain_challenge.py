"""Tests for the TDBRAIN-Challenge template filler (real V3.1 MDD-Dx layout).

The real replication template (``TDBRAIN_MDD_Dx_Prediction_replicationV3.xlsx``)
is one sheet ``Blad1`` with header
``["ID", "Diagnosis (MDD=1; Control=0)", "Consent", "age", "gender"]`` and 60
rows (ID 1..60), the three prediction columns pre-filled with the placeholder
string ``"REPLICATION"``. A submission must overwrite ONLY those three columns
(diagnosis, age, gender) for the right subject, leaving ID/Consent untouched.
"""

from __future__ import annotations

import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))
import numpy as np  # noqa: E402

from tdbrain_challenge import _apply_norm, _resolve_column, fill_challenge_template  # noqa: E402

REAL_HEADER = ["ID", "Diagnosis (MDD=1; Control=0)", "Consent", "age", "gender"]


def _write_template(tmp_path, n=3):
    """Synthesize a copy of the real MDD-Dx replication template."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Blad1"
    ws.append(REAL_HEADER)
    for i in range(1, n + 1):
        ws.append([i, "REPLICATION", "YES", "REPLICATION", "REPLICATION"])
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return str(p)


def test_resolve_column_maps_targets_to_real_headers():
    assert _resolve_column(REAL_HEADER, "diagnosis") == 1
    assert _resolve_column(REAL_HEADER, "age") == 3
    assert _resolve_column(REAL_HEADER, "gender") == 4
    assert _resolve_column(REAL_HEADER, "id") == 0


def test_fill_challenge_template_writes_three_targets(tmp_path):
    template = _write_template(tmp_path, n=3)
    out = str(tmp_path / "submission.xlsx")
    preds = {
        "diagnosis": {1: 1, 2: 0, 3: 1},
        "age": {1: 41.2, 2: 33.7, 3: 58.0},
        "gender": {1: 0, 2: 1, 3: 0},
    }
    fill_challenge_template(template, out, preds)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    # diagnosis / gender written as ints, age as float; Consent + ID untouched
    assert rows[1][1].value == 1 and rows[2][1].value == 0 and rows[3][1].value == 1
    assert abs(rows[1][3].value - 41.2) < 1e-6
    assert rows[1][4].value == 0 and rows[2][4].value == 1
    assert rows[1][2].value == "YES"  # Consent preserved
    # no placeholder strings survive in the three prediction columns
    for r in ws.iter_rows(min_row=2):
        assert r[1].value != "REPLICATION"
        assert r[3].value != "REPLICATION"
        assert r[4].value != "REPLICATION"


def test_apply_norm_recording_zscore_is_affine_invariant():
    # BrainVision-vs-BDF export differs mainly by a per-recording affine
    # (gain+contrast) transform of the log-spectrum; per-recording z-score of the
    # band block must be invariant to it so a Discovery-trained model transfers.
    rng = np.random.RandomState(0)
    base = rng.randn(4, 156)
    shifted = base.copy()
    shifted[:, :130] = base[:, :130] * 2.0 + 5.0  # affine on band block only
    zb = _apply_norm(base, "recording-zscore")
    zs = _apply_norm(shifted, "recording-zscore")
    assert np.allclose(zb[:, :130], zs[:, :130], atol=1e-6)      # band: affine-invariant
    assert np.allclose(zb[:, 130:], base[:, 130:])                # slope: untouched
    assert np.allclose(zb[:, :130].mean(1), 0, atol=1e-6)
    assert np.allclose(zb[:, :130].std(1), 1, atol=1e-3)


def test_apply_norm_none_is_identity():
    x = np.arange(156, dtype=float).reshape(1, 156)
    assert np.array_equal(_apply_norm(x, "none"), x)
    assert np.array_equal(_apply_norm(x, None), x)


def test_fill_challenge_template_skips_missing_predictions(tmp_path):
    # a subject with no prediction (e.g. unloadable EEG) keeps the placeholder,
    # so a partial submission is visibly incomplete rather than silently 0.
    template = _write_template(tmp_path, n=3)
    out = str(tmp_path / "submission.xlsx")
    preds = {"diagnosis": {1: 1, 3: 0}, "age": {}, "gender": {}}
    fill_challenge_template(template, out, preds)
    wb = openpyxl.load_workbook(out)
    rows = {r[0].value: r for r in wb.active.iter_rows(min_row=2)}
    assert rows[1][1].value == 1
    assert rows[2][1].value == "REPLICATION"  # no prediction -> untouched
    assert rows[3][1].value == 0
